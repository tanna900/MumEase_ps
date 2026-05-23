from fastapi import APIRouter, Request, UploadFile, File, Form, HTTPException
from fastapi.responses import StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from typing import List, Dict, Tuple
import csv, io, re, os

from openpyxl import Workbook, load_workbook

from app.database import SessionLocal
from app import models

router = APIRouter(prefix="/products", tags=["Products Import"])
templates = Jinja2Templates(directory="app/templates")

REQUIRED_COLS = ["name", "barcode"]
OPTIONAL_COLS = ["color", "size", "cost_price", "price", "stock"]
ALL_COLS = REQUIRED_COLS + OPTIONAL_COLS


# ----------------- DB utils -----------------
def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


# ----------------- Helpers -----------------
def _to_float(v):
    try:
        if v is None or str(v).strip() == "":
            return None
        return float(str(v).replace(",", "").strip())
    except:
        return None

def _to_int(v):
    try:
        if v is None or str(v).strip() == "":
            return None
        return int(float(str(v).replace(",", "").strip()))
    except:
        return None

def _normalize_headers(headers: List[str]) -> List[str]:
    """
    نحاول نطابق أعمدة الإكسل/CSV لأسماءنا القياسية بغض النظر عن المسافات/حالة الأحرف.
    """
    norm = []
    for h in headers:
        k = (h or "").strip().lower()
        k = re.sub(r"\s+", "_", k)
        norm.append(k)
    return norm

def _read_csv_to_rows(raw_bytes: bytes) -> Tuple[List[str], List[Dict]]:
    text = raw_bytes.decode("utf-8-sig", errors="ignore")
    f = io.StringIO(text)
    reader = csv.DictReader(f)
    headers = _normalize_headers(reader.fieldnames or [])
    rows = []
    for row in reader:
        # طَبِّق نفس التطبيع على المفاتيح
        r = { (k or "").strip().lower().replace(" ", "_"): (v or "").strip() for k, v in row.items() }
        rows.append(r)
    return headers, rows

def _read_xlsx_to_rows(raw_bytes: bytes) -> Tuple[List[str], List[Dict]]:
    f = io.BytesIO(raw_bytes)
    wb = load_workbook(f, data_only=True)
    ws = wb.active  # أول شيت

    # الهيدر من أول صف
    headers_raw = []
    for cell in ws[1]:
        headers_raw.append((cell.value or "").strip() if isinstance(cell.value, str) else (cell.value or ""))

    headers = _normalize_headers([str(h) for h in headers_raw])

    rows: List[Dict] = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        rdict: Dict = {}
        for i, v in enumerate(row):
            key = headers[i] if i < len(headers) else f"col{i+1}"
            value = "" if v is None else str(v).strip() if isinstance(v, str) else v
            rdict[key] = value
        # تجاهل الصفوف الفارغة تمامًا
        if any(str(v).strip() for v in rdict.values()):
            rows.append(rdict)
    return headers, rows

def _validate_required(headers: List[str]):
    missing = [c for c in REQUIRED_COLS if c not in headers]
    if missing:
        raise HTTPException(status_code=400, detail=f"أعمدة ناقصة: {', '.join(missing)}")

def _row_to_product_fields(row: Dict) -> Dict:
    """
    يحوّل صف (dict) لقيم الحقول القياسية عندنا.
    بيقبل أسماء أعمدة مطبّعة (lower/underscored).
    """
    name = (row.get("name") or "").strip()
    barcode = (row.get("barcode") or "").strip()

    color = (row.get("color") or None) or None
    size = (row.get("size") or None) or None
    cost_price = _to_float(row.get("cost_price"))
    price = _to_float(row.get("price"))
    stock = _to_int(row.get("stock"))

    return {
        "name": name,
        "barcode": barcode,
        "color": color or None,
        "size": size or None,
        "cost_price": cost_price,
        "price": price,
        "stock": stock,
    }


# ----------------- Endpoints -----------------
@router.get("/import")
def import_form(request: Request):
    """
    صفحة الاستيراد: ترفع ملف XLSX/CSV + اختيار نمط الاستيراد (إضافة فقط أو تحديث/إضافة)
    """
    return templates.TemplateResponse("products_import.html", {"request": request})


@router.get("/import/template.xlsx")
def download_template_xlsx():
    """
    تنزيل قالب Excel جاهز بالأعمدة الصحيحة.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(ALL_COLS)
    # مثال صف اختياري
    ws.append(["T-Shirt Nursing", "1234567890123", "Black", "L", "120", "199.99", "50"])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="products_template.xlsx"'},
    )


@router.get("/import/template.csv")
def download_template_csv():
    """
    تنزيل قالب CSV جاهز بالأعمدة الصحيحة.
    """
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(ALL_COLS)
    writer.writerow(["T-Shirt Nursing", "1234567890123", "Black", "L", "120", "199.99", "50"])
    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="products_template.csv"'},
    )


@router.post("/import")
async def import_file(
    request: Request,
    file: UploadFile = File(...),
    mode: str = Form("upsert"),  # upsert أو insert
):
    """
    استيراد XLSX أو CSV:
    - الأعمدة المطلوبة: name, barcode
    - المسموح: color, size, cost_price, price, stock
    - mode:
        * insert: يضيف الجديد فقط، ولو باركود موجود -> يتخطى
        * upsert: يضيف الجديد ويحدّث الموجود (حسب barcode)
    """
    raw = await file.read()
    filename = (file.filename or "").lower()
    is_xlsx = filename.endswith(".xlsx")
    is_csv = filename.endswith(".csv")

    if not (is_xlsx or is_csv):
        raise HTTPException(status_code=400, detail="صيغة غير مدعومة. من فضلك ارفع ملف XLSX أو CSV.")

    # اقرأ الملف وأرجع headers + rows
    if is_xlsx:
        headers, rows = _read_xlsx_to_rows(raw)
    else:
        headers, rows = _read_csv_to_rows(raw)

    _validate_required(headers)

    db: Session = next(get_db())
    created, updated, skipped = 0, 0, 0
    errors: List[Dict] = []

    # نبدأ من 2 في الإكسل (الهيدر سطر 1) — بس هنستخدم ترقيم متّسق للجميع
    current_line = 1  # الهيدر
    for row in rows:
        current_line += 1
        try:
            fields = _row_to_product_fields(row)
            name = fields["name"]
            barcode = fields["barcode"]

            if not name or not barcode:
                skipped += 1
                errors.append({"line": current_line, "reason": "name/barcode فارغ"})
                continue

            existing = db.query(models.Product).filter(models.Product.barcode == barcode).first()
            if existing:
                if mode == "insert":
                    skipped += 1
                    continue
                # تحديث
                existing.name = name
                existing.color = fields["color"]
                existing.size = fields["size"]
                if fields["cost_price"] is not None: existing.cost_price = fields["cost_price"]
                if fields["price"] is not None:      existing.price = fields["price"]
                if fields["stock"] is not None:      existing.stock = fields["stock"]
                updated += 1
            else:
                p = models.Product(
                    name=name,
                    barcode=barcode,
                    color=fields["color"] or None,
                    size=fields["size"] or None,
                    cost_price=fields["cost_price"] or 0,
                    price=fields["price"] or 0,
                    stock=fields["stock"] or 0,
                )
                db.add(p)
                created += 1

        except Exception as ex:
            errors.append({"line": current_line, "reason": str(ex)})
            skipped += 1

    db.commit()

    ctx = {
        "request": request,
        "result": {
            "created": created,
            "updated": updated,
            "skipped": skipped,
            "errors": errors,
            "mode": mode,
            "filename": file.filename,
            "rows_count": len(rows),
        }
    }
    return templates.TemplateResponse("products_import.html", ctx)