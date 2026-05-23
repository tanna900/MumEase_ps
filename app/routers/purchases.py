# app/routers/purchases.py
from fastapi import APIRouter, Request, Depends, Form, UploadFile, File
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import desc, func, text, inspect
from app.database import SessionLocal
from app import models
from fastapi.templating import Jinja2Templates

from datetime import datetime
import io
import openpyxl

import os
import shutil
from datetime import datetime as _dt

router = APIRouter(prefix="/purchases", tags=["Purchases"])
templates = Jinja2Templates(directory="app/templates")

# ---------------- DB ----------------
def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

def _back_url(msg: str = ""):
    if msg:
        return f"/purchases?msg={msg}"
    return "/purchases"

# =========================
# ✅ AUTO MIGRATION (No terminal needed)
# =========================
_PURCHASE_MIGRATED = False

def _backup_sqlite_if_possible(engine):
    try:
        if engine.dialect.name != "sqlite":
            return
        db_path = str(engine.url.database or "")
        if db_path and os.path.exists(db_path):
            ts = _dt.now().strftime("%Y%m%d_%H%M%S")
            backup_path = f"{db_path}.backup_{ts}"
            shutil.copy2(db_path, backup_path)
    except Exception:
        pass

def _ensure_purchase_schema(db: Session):
    """
    يضيف الأعمدة الجديدة في purchase_items تلقائيًا لو مش موجودة.
    """
    global _PURCHASE_MIGRATED
    if _PURCHASE_MIGRATED:
        return

    try:
        engine = db.get_bind()
        inspector = inspect(engine)
        tables = set(inspector.get_table_names())
        if "purchase_items" not in tables:
            _PURCHASE_MIGRATED = True
            return

        _backup_sqlite_if_possible(engine)

        existing_cols = {c["name"] for c in inspector.get_columns("purchase_items")}
        dialect = engine.dialect.name

        def col_float():
            if dialect == "postgresql":
                return "DOUBLE PRECISION"
            if dialect == "mysql":
                return "DOUBLE"
            return "REAL"

        def col_int():
            if dialect == "postgresql":
                return "INTEGER"
            if dialect == "mysql":
                return "INT"
            return "INTEGER"

        alters = []

        # ✅ الأعمدة الموجودة سابقًا
        if "unit_mfg_cost" not in existing_cols:
            alters.append(f"ALTER TABLE purchase_items ADD COLUMN unit_mfg_cost {col_float()} DEFAULT 0")
        if "unit_total_cost" not in existing_cols:
            alters.append(f"ALTER TABLE purchase_items ADD COLUMN unit_total_cost {col_float()} DEFAULT 0")

        # ✅ جديد: خامات + نثريات
        if "unit_fabric_cost" not in existing_cols:
            alters.append(f"ALTER TABLE purchase_items ADD COLUMN unit_fabric_cost {col_float()} DEFAULT 0")
        if "unit_overhead_cost" not in existing_cols:
            alters.append(f"ALTER TABLE purchase_items ADD COLUMN unit_overhead_cost {col_float()} DEFAULT 0")

        # ✅ rollback الآمن
        if "old_stock" not in existing_cols:
            alters.append(f"ALTER TABLE purchase_items ADD COLUMN old_stock {col_int()} DEFAULT 0")
        if "applied_cost_price" not in existing_cols:
            alters.append(f"ALTER TABLE purchase_items ADD COLUMN applied_cost_price {col_float()} DEFAULT 0")
        if "applied_price" not in existing_cols:
            alters.append(f"ALTER TABLE purchase_items ADD COLUMN applied_price {col_float()} DEFAULT 0")
        if "applied_stock" not in existing_cols:
            alters.append(f"ALTER TABLE purchase_items ADD COLUMN applied_stock {col_int()} DEFAULT 0")

        for sql in alters:
            try:
                db.execute(text(sql))
            except Exception:
                pass

        # ✅ تهيئة الداتا القديمة: لو unit_total_cost فاضي خليها = unit_cost القديم
        try:
            db.execute(text("""
                UPDATE purchase_items
                SET unit_total_cost = unit_cost
                WHERE (unit_total_cost IS NULL OR unit_total_cost = 0)
                  AND unit_cost IS NOT NULL
                  AND unit_cost <> 0
            """))
        except Exception:
            pass

        # ✅ Indexes (SQLite فقط بشكل آمن)
        try:
            if dialect == "sqlite":
                db.execute(text("CREATE INDEX IF NOT EXISTS ix_purchase_items_purchase_id ON purchase_items(purchase_id)"))
                db.execute(text("CREATE INDEX IF NOT EXISTS ix_purchase_items_product_id ON purchase_items(product_id)"))
                db.execute(text("CREATE INDEX IF NOT EXISTS ix_purchase_items_barcode ON purchase_items(barcode)"))
        except Exception:
            pass

        db.commit()
        _PURCHASE_MIGRATED = True

    except Exception:
        try:
            db.rollback()
        except Exception:
            pass
        _PURCHASE_MIGRATED = True

# ---------- Helpers ----------
def _safe_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return str(v).strip()

def _parse_factory_id(raw: str):
    s = _safe_str(raw)
    if not s:
        return None
    try:
        x = int(s)
        return x if x > 0 else None
    except Exception:
        return None

def _avg_cost(total_qty: int, total_cost: float) -> float:
    if total_qty and total_qty > 0:
        return float(total_cost) / float(total_qty)
    return 0.0

def _weighted_avg(old_stock: int, old_avg: float, add_qty: int, add_unit_cost: float) -> float:
    new_stock = old_stock + add_qty
    if new_stock <= 0:
        return 0.0
    return ((old_stock * old_avg) + (add_qty * add_unit_cost)) / new_stock

# ========== توليد رقم فاتورة شراء ==========
def _generate_po_code(db: Session) -> str:
    today = datetime.now().strftime("%Y%m%d")
    base = f"PO-{today}-"
    count = (
        db.query(func.count(models.PurchaseInvoice.id))
        .filter(models.PurchaseInvoice.code.like(f"{base}%"))
        .scalar()
        or 0
    )
    return f"{base}{count+1:04d}"

# ---------- Excel parsing ----------
def _ws_is_header_format(ws) -> bool:
    try:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    except StopIteration:
        return False
    hdrs = [(_safe_str(x).lower()) for x in (first or [])]
    if not (("barcode" in hdrs) and ("qty" in hdrs or "quantity" in hdrs)):
        return False
    # نقبل: total_cost أو cost_price أو fabric_cost (لو عايز تحسب total تلقائي)
    return ("total_cost" in hdrs) or ("cost_price" in hdrs) or ("fabric_cost" in hdrs)

def _iter_rows_any_format(ws):
    """
    يرجّع صفوف موحّدة (10 حقول):
    barcode, name, color, size, qty, unit_fabric_cost, unit_mfg_cost, unit_overhead_cost, unit_total_cost, unit_price

    - Headers format (الجديد):
        barcode, name, color, size, qty, fabric_cost, mfg_cost, overhead_cost, total_cost, price
      * total_cost اختياري لو هتسيبه يتحسب = fabric + mfg + overhead
    - Compatible:
        cost_price بدل total_cost (نعتبره total_cost)
    - Old A..G:
        A Barcode, B Name, C Color, D Size, E Qty, F Cost (total_cost), G Price
        fabric/mfg/overhead = 0
    """
    if _ws_is_header_format(ws):
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        hdrs = [(_safe_str(x).lower()) for x in (header or [])]
        idx = {h: i for i, h in enumerate(hdrs) if h}

        def get(row, key, default=""):
            i = idx.get(key)
            if i is None or i >= len(row):
                return default
            return row[i]

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue

            barcode = _safe_str(get(row, "barcode"))
            name = _safe_str(get(row, "name"))
            color = _safe_str(get(row, "color"))
            size = _safe_str(get(row, "size"))
            qty = int(get(row, "qty", 0) or 0)

            unit_fabric_cost = float(get(row, "fabric_cost", 0) or 0)
            unit_mfg_cost = float(get(row, "mfg_cost", 0) or 0)
            unit_overhead_cost = float(get(row, "overhead_cost", 0) or 0)

            # total_cost: لو موجود خدّه، لو مش موجود احسبه
            if "total_cost" in idx:
                unit_total_cost = float(get(row, "total_cost", 0) or 0)
            elif "cost_price" in idx:
                unit_total_cost = float(get(row, "cost_price", 0) or 0)
            else:
                unit_total_cost = float(unit_fabric_cost + unit_mfg_cost + unit_overhead_cost)

            unit_price = float(get(row, "price", 0) or 0)

            yield barcode, name, color, size, qty, unit_fabric_cost, unit_mfg_cost, unit_overhead_cost, unit_total_cost, unit_price

    else:
        # Old A..G
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            barcode = _safe_str(row[0] if len(row) > 0 else "")
            name = _safe_str(row[1] if len(row) > 1 else "")
            color = _safe_str(row[2] if len(row) > 2 else "")
            size = _safe_str(row[3] if len(row) > 3 else "")
            qty = int((row[4] if len(row) > 4 else 0) or 0)

            unit_total_cost = float((row[5] if len(row) > 5 else 0) or 0)
            unit_price = float((row[6] if len(row) > 6 else 0) or 0)

            unit_fabric_cost = 0.0
            unit_mfg_cost = 0.0
            unit_overhead_cost = 0.0

            yield barcode, name, color, size, qty, unit_fabric_cost, unit_mfg_cost, unit_overhead_cost, unit_total_cost, unit_price

# ---------- Notes ----------
def _po_total_note(code: str, pid: int, action: str) -> str:
    if action == "add":
        return f"[PO {code} PID={pid}] تصنيع — إجمالي الفاتورة"
    if action == "edit_rev":
        return f"[ADJ][PO {code} PID={pid}] تعديل — إلغاء تصنيع القديم (إجمالي)"
    if action == "edit_add":
        return f"[PO {code} PID={pid}] تعديل — إضافة تصنيع الجديد (إجمالي)"
    if action == "delete":
        return f"[ADJ][PO {code} PID={pid}] إلغاء تصنيع (إجمالي)"
    return f"[PO {code} PID={pid}] تصنيع (إجمالي)"

def _pm_note(code: str, action: str, qty: int,
             unit_total: float = 0.0, unit_mfg: float = 0.0, unit_fabric: float = 0.0, unit_over: float = 0.0, unit_price: float = 0.0) -> str:
    if action == "add":
        return f"[PO {code}] +{qty} | fabric={unit_fabric:.2f} | mfg={unit_mfg:.2f} | over={unit_over:.2f} | total={unit_total:.2f} | sell={unit_price:.2f}"
    if action == "revert":
        return f"[ADJ][PO {code}] -{abs(qty)} | تعديل مشتريات (إلغاء القديم)"
    if action == "apply":
        return f"[PO {code}] +{qty} | fabric={unit_fabric:.2f} | mfg={unit_mfg:.2f} | over={unit_over:.2f} | total={unit_total:.2f} | sell={unit_price:.2f} | تعديل مشتريات (الجديد)"
    if action == "delete":
        return f"[ADJ][PO {code}] -{abs(qty)} | إلغاء إذن مشتريات"
    return f"[PO {code}] {qty}"

def _calc_mfg_totals(purchase: models.PurchaseInvoice):
    total_qty = 0
    total_mfg = 0.0
    for it in purchase.items:
        q = int(it.qty or 0)
        if q <= 0:
            continue
        m = float(getattr(it, "unit_mfg_cost", 0) or 0)
        total_qty += q
        total_mfg += (q * m)
    avg_mfg = _avg_cost(total_qty, total_mfg)
    return total_qty, total_mfg, avg_mfg

def _calc_fabric_total_from_items(purchase: models.PurchaseInvoice) -> float:
    total = 0.0
    for it in purchase.items:
        q = int(it.qty or 0)
        if q <= 0:
            continue
        f = float(getattr(it, "unit_fabric_cost", 0) or 0)
        total += q * f
    return round(total, 2)

# =========================
# خامات انتاج: نضمن الجدول موجود + نسجل حركة OUT/IN
# =========================
_MAT_ENSURED = False

def _ensure_materials_schema(db: Session):
    global _MAT_ENSURED
    if _MAT_ENSURED:
        return
    try:
        engine = db.get_bind()
        inspector = inspect(engine)
        tables = set(inspector.get_table_names())
        if "production_material_entries" not in tables:
            db.execute(text("""
            CREATE TABLE IF NOT EXISTS production_material_entries (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                date TEXT,
                type TEXT,
                amount REAL DEFAULT 0,
                note TEXT,
                finance_entry_id INTEGER,
                ref_type TEXT,
                ref_id INTEGER,
                created_at TEXT
            )
            """))
            db.execute(text("CREATE INDEX IF NOT EXISTS ix_pme_date ON production_material_entries(date)"))
            db.execute(text("CREATE INDEX IF NOT EXISTS ix_pme_type ON production_material_entries(type)"))
            db.execute(text("CREATE INDEX IF NOT EXISTS ix_pme_ref ON production_material_entries(ref_type, ref_id)"))
            db.execute(text("CREATE INDEX IF NOT EXISTS ix_pme_fin ON production_material_entries(finance_entry_id)"))
            db.commit()
        _MAT_ENSURED = True
    except Exception:
        try:
            db.rollback()
        except Exception:
            pass
        _MAT_ENSURED = True

def _materials_post(db: Session, type_: str, amount: float, date: str, note: str, ref_type: str, ref_id: int):
    """
    يسجل حركة في خامات انتاج فقط (بدون خزنة).
    OUT = خصم خامات / IN = رجوع خامات.
    """
    if amount <= 0:
        return
    _ensure_materials_schema(db)

    # لو ORM model موجود
    if hasattr(models, "ProductionMaterialEntry"):
        db.add(models.ProductionMaterialEntry(
            date=date,
            type=type_,
            amount=float(amount),
            note=(note or "").strip() or None,
            finance_entry_id=None,
            ref_type=ref_type,
            ref_id=int(ref_id),
        ))
        return

    # fallback SQL
    db.execute(text("""
        INSERT INTO production_material_entries (date, type, amount, note, finance_entry_id, ref_type, ref_id, created_at)
        VALUES (:d, :t, :a, :n, NULL, :rt, :rid, :c)
    """), {
        "d": date,
        "t": type_,
        "a": float(amount),
        "n": (note or "").strip() or None,
        "rt": ref_type,
        "rid": int(ref_id),
        "c": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    })

# =========================
# تنزيل قالب Excel
# =========================
@router.get("/import/template.xlsx")
def download_purchase_template(db: Session = Depends(get_db)):
    _ensure_purchase_schema(db)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "purchases"

    headers = ["barcode", "name", "color", "size", "qty", "fabric_cost", "mfg_cost", "overhead_cost", "total_cost", "price"]
    ws.append(headers)
    # مثال
    # total_cost = fabric + mfg + overhead
    ws.append(["1234567890", "Hoodie", "Black", "L", 10, 80, 20, 5, 105, 299])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="purchase_template.xlsx"'},
    )

# ========== صفحة عرض فواتير المشتريات ==========
@router.get("", response_class=HTMLResponse)
def list_purchases(request: Request, msg: str = "", db: Session = Depends(get_db)):
    _ensure_purchase_schema(db)

    purchases = db.query(models.PurchaseInvoice).order_by(desc(models.PurchaseInvoice.id)).all()
    factories = db.query(models.Factory).all()
    factories_map = {f.id: f.name for f in factories}

    return templates.TemplateResponse(
        "purchases_list.html",
        {"request": request, "msg": msg, "purchases": purchases, "factories_map": factories_map},
    )

@router.get("/import", response_class=HTMLResponse)
def import_purchases_page(request: Request, msg: str = "", db: Session = Depends(get_db)):
    _ensure_purchase_schema(db)

    purchases = db.query(models.PurchaseInvoice).order_by(desc(models.PurchaseInvoice.id)).all()
    factories = db.query(models.Factory).all()
    factories_map = {f.id: f.name for f in factories}

    return templates.TemplateResponse(
        "purchases_list.html",
        {"request": request, "msg": msg, "purchases": purchases, "factories_map": factories_map},
    )

# ========== رفع شيت مشتريات جديد ==========
@router.post("/upload")
async def upload_purchase(
    file: UploadFile = File(...),
    date: str = Form(""),
    factory_id: str = Form(""),
    note: str = Form(""),
    code: str = Form(""),
    db: Session = Depends(get_db),
):
    _ensure_purchase_schema(db)

    if not date:
        date = datetime.now().strftime("%Y-%m-%d")

    code = (code or "").strip()
    if not code:
        code = _generate_po_code(db)

    factory_id_int = _parse_factory_id(factory_id)

    try:
        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
    except Exception:
        db.rollback()
        return RedirectResponse(url=_back_url("⚠️ ملف الإكسيل غير صالح"), status_code=303)

    try:
        purchase = models.PurchaseInvoice(
            code=code,
            date=date,
            factory_id=factory_id_int,
            total_qty=0,
            total_cost=0,
            file_name=file.filename,
            note=note,
            history=f"[{datetime.now().strftime('%Y-%m-%d %H:%M')}] إنشاء الفاتورة من ملف: {file.filename}",
        )
        db.add(purchase)
        db.flush()

        total_qty = 0
        total_cost = 0.0
        total_mfg_cost = 0.0
        total_fabric_cost = 0.0
        total_overhead_cost = 0.0

        for barcode, name, color, size, qty, unit_fabric_cost, unit_mfg_cost, unit_over_cost, unit_total_cost, unit_price in _iter_rows_any_format(ws):
            if not barcode or qty <= 0:
                continue

            unit_fabric_cost = float(unit_fabric_cost or 0)
            unit_mfg_cost = float(unit_mfg_cost or 0)
            unit_over_cost = float(unit_over_cost or 0)
            unit_total_cost = float(unit_total_cost or 0)
            unit_price = float(unit_price or 0)

            # لو total_cost متساب 0 في الشيت الجديد، نحسبه تلقائيًا
            if unit_total_cost <= 0 and (unit_fabric_cost + unit_mfg_cost + unit_over_cost) > 0:
                unit_total_cost = float(unit_fabric_cost + unit_mfg_cost + unit_over_cost)

            # شرط منطقي
            if unit_total_cost + 1e-9 < (unit_fabric_cost + unit_mfg_cost + unit_over_cost):
                db.rollback()
                return RedirectResponse(
                    url=_back_url("⚠️ خطأ: total_cost لازم يكون >= (fabric_cost + mfg_cost + overhead_cost)"),
                    status_code=303
                )

            product = db.query(models.Product).filter(models.Product.barcode == barcode).first()
            if not product:
                product = models.Product(
                    name=name or barcode,
                    barcode=barcode,
                    color=color or None,
                    size=size or None,
                    cost_price=0,
                    price=0,
                    stock=0,
                )
                db.add(product)
                db.flush()

            prev_cost_avg = float(product.cost_price or 0)
            prev_price = float(product.price or 0)
            prev_stock = int(product.stock or 0)

            new_stock = prev_stock + int(qty)
            new_cost_avg = _weighted_avg(prev_stock, prev_cost_avg, int(qty), unit_total_cost)

            new_price = prev_price
            if unit_price > 0:
                new_price = unit_price

            product.stock = new_stock
            product.cost_price = new_cost_avg
            product.price = new_price

            item = models.PurchaseItem(
                purchase_id=purchase.id,
                product_id=product.id,
                barcode=barcode,
                name=name or product.name,
                color=color or product.color,
                size=size or product.size,
                qty=int(qty),

                unit_fabric_cost=unit_fabric_cost,
                unit_mfg_cost=unit_mfg_cost,
                unit_overhead_cost=unit_over_cost,
                unit_total_cost=unit_total_cost,

                unit_cost=unit_total_cost,  # legacy
                unit_price=unit_price,

                old_cost_price=prev_cost_avg,
                old_price=prev_price,
                old_stock=prev_stock,

                applied_cost_price=new_cost_avg,
                applied_price=new_price,
                applied_stock=new_stock,
            )
            db.add(item)

            if hasattr(models, "ProductMovement"):
                mv = models.ProductMovement(
                    product_id=product.id,
                    movement_type="purchase",
                    qty=int(qty),
                    balance_after=int(product.stock or 0),
                    ref_type="PurchaseInvoice",
                    ref_id=int(purchase.id),
                    note=_pm_note(code, "add", int(qty), unit_total_cost, unit_mfg_cost, unit_fabric_cost, unit_over_cost, unit_price),
                )
                db.add(mv)

            total_qty += int(qty)
            total_cost += int(qty) * unit_total_cost
            total_mfg_cost += int(qty) * unit_mfg_cost
            total_fabric_cost += int(qty) * unit_fabric_cost
            total_overhead_cost += int(qty) * unit_over_cost

        purchase.total_qty = float(total_qty)
        purchase.total_cost = float(total_cost)

        # ✅ حساب المصنع: سطر واحد إجمالي التصنيع
        if purchase.factory_id and hasattr(models, "ManufacturingBatch") and total_qty > 0:
            avg_mfg = _avg_cost(int(total_qty), float(total_mfg_cost))
            mb = models.ManufacturingBatch(
                date=date,
                factory_id=purchase.factory_id,
                product_id=None,
                qty=int(total_qty),
                unit_cost=float(avg_mfg),
                note=_po_total_note(code, purchase.id, "add"),
            )
            db.add(mb)

        # ✅ خامات انتاج: خصم "القماش فقط" مرة واحدة
        if total_fabric_cost > 0:
            _materials_post(
                db,
                type_="OUT",
                amount=round(float(total_fabric_cost), 2),
                date=date,
                note=f"[PO {code}] خصم خامات (قماش فقط) من فاتورة المشتريات",
                ref_type="PurchaseInvoice",
                ref_id=int(purchase.id),
            )

        db.commit()
        return RedirectResponse(url=_back_url(f"✅ تم تسجيل فاتورة مشتريات {code} بعدد {total_qty} قطعة"), status_code=303)

    except Exception:
        db.rollback()
        return RedirectResponse(url=_back_url("⚠️ حصل خطأ أثناء تسجيل المشتريات"), status_code=303)

# ========== عرض تفاصيل فاتورة مشتريات ==========
@router.get("/{pid}", response_class=HTMLResponse)
def purchase_detail(request: Request, pid: int, msg: str = "", db: Session = Depends(get_db)):
    _ensure_purchase_schema(db)

    purchase = db.query(models.PurchaseInvoice).get(pid)
    if not purchase:
        return RedirectResponse(url=_back_url("⚠️ الفاتورة غير موجودة"), status_code=303)

    factory_name = "-"
    if purchase.factory_id:
        f = db.query(models.Factory).get(purchase.factory_id)
        if f:
            factory_name = f.name

    factories = db.query(models.Factory).all()
    factories_map = {f.id: f.name for f in factories}

    return templates.TemplateResponse(
        "purchase_detail.html",
        {
            "request": request,
            "purchase": purchase,
            "factory_name": factory_name,
            "factories_map": factories_map,
            "msg": msg,
        },
    )

# ========== حذف فاتورة مشتريات ==========
@router.post("/{pid}/delete")
def delete_purchase(pid: int, db: Session = Depends(get_db)):
    _ensure_purchase_schema(db)

    purchase = db.query(models.PurchaseInvoice).get(pid)
    if not purchase:
        return RedirectResponse(url=_back_url("⚠️ الفاتورة غير موجودة"), status_code=303)

    try:
        code = purchase.code or str(purchase.id)
        p_date = purchase.date or datetime.now().strftime("%Y-%m-%d")

        # ✅ رجوع خامات انتاج (قماش فقط) قبل ما نمسح البنود
        old_fabric_total = _calc_fabric_total_from_items(purchase)
        if old_fabric_total > 0:
            _materials_post(
                db,
                type_="IN",
                amount=float(old_fabric_total),
                date=p_date,
                note=f"[PO {code}] رجوع خامات (قماش فقط) بسبب حذف فاتورة المشتريات",
                ref_type="PurchaseInvoice",
                ref_id=int(purchase.id),
            )

        # ✅ حساب المصنع: قيد عكسي واحد إجمالي التصنيع
        if purchase.factory_id and hasattr(models, "ManufacturingBatch") and (purchase.total_qty or 0) > 0:
            total_qty, total_mfg_cost, avg_mfg = _calc_mfg_totals(purchase)
            if total_qty > 0:
                mb = models.ManufacturingBatch(
                    date=p_date,
                    factory_id=purchase.factory_id,
                    product_id=None,
                    qty=-(int(total_qty)),
                    unit_cost=float(avg_mfg),
                    note=_po_total_note(code, purchase.id, "delete"),
                )
                db.add(mb)

        # ✅ رجوع المخزون + رجوع cost/price بشكل "آمن"
        for item in purchase.items:
            product = db.query(models.Product).get(item.product_id)
            if not product:
                continue

            dec_qty = int(item.qty or 0)
            if dec_qty <= 0:
                continue

            old_stock_now = int(product.stock or 0)
            product.stock = max(0, old_stock_now - dec_qty)

            try:
                if (
                    int(product.stock + dec_qty) == int(getattr(item, "applied_stock", 0) or 0)
                    and float(product.cost_price or 0) == float(getattr(item, "applied_cost_price", 0) or 0)
                    and float(product.price or 0) == float(getattr(item, "applied_price", 0) or 0)
                ):
                    product.cost_price = float(getattr(item, "old_cost_price", 0) or 0)
                    product.price = float(getattr(item, "old_price", 0) or 0)
            except Exception:
                pass

            if hasattr(models, "ProductMovement"):
                mv = models.ProductMovement(
                    product_id=product.id,
                    movement_type="purchase_delete",
                    qty=-(dec_qty),
                    balance_after=int(product.stock or 0),
                    ref_type="PurchaseInvoice",
                    ref_id=int(purchase.id),
                    note=_pm_note(code, "delete", -dec_qty),
                )
                db.add(mv)

        now = datetime.now().strftime("%Y-%m-%d %H:%M")
        h = purchase.history or ""
        h += f"\n[{now}] تم حذف الفاتورة وإرجاع المخزون + قيد تصنيع عكسي + رجوع خامات (قماش فقط)."
        purchase.history = h

        db.delete(purchase)
        db.commit()
        return RedirectResponse(url=_back_url("🗑️ تم حذف فاتورة المشتريات وإرجاع المخزون + خامات القماش"), status_code=303)

    except Exception:
        db.rollback()
        return RedirectResponse(url=_back_url("⚠️ حصل خطأ أثناء حذف الفاتورة"), status_code=303)

# ========== تعديل فاتورة (رفع شيت جديد مكان القديم) ==========
@router.post("/{pid}/edit")
async def edit_purchase(
    pid: int,
    file: UploadFile = File(...),
    date: str = Form(""),
    factory_id: str = Form(""),
    note: str = Form(""),
    db: Session = Depends(get_db),
):
    _ensure_purchase_schema(db)

    purchase = db.query(models.PurchaseInvoice).get(pid)
    if not purchase:
        return RedirectResponse(url=_back_url("⚠️ الفاتورة غير موجودة"), status_code=303)

    try:
        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
    except Exception:
        db.rollback()
        return RedirectResponse(url=_back_url("⚠️ ملف الإكسيل الجديد غير صالح"), status_code=303)

    try:
        if not date:
            date = purchase.date or datetime.now().strftime("%Y-%m-%d")

        code = purchase.code or str(purchase.id)

        old_factory_id = purchase.factory_id
        new_factory_id = old_factory_id

        factory_id_int = _parse_factory_id(factory_id)
        if factory_id_int:
            new_factory_id = factory_id_int
            purchase.factory_id = factory_id_int

        # ✅ 0) رجوع خامات القماش القديمة (قبل مسح البنود)
        old_fabric_total = _calc_fabric_total_from_items(purchase)
        if old_fabric_total > 0:
            _materials_post(
                db,
                type_="IN",
                amount=float(old_fabric_total),
                date=purchase.date or date,
                note=f"[PO {code}] تعديل — رجوع خامات (قماش فقط) للقديم",
                ref_type="PurchaseInvoice",
                ref_id=int(purchase.id),
            )

        # ✅ 1) حساب المصنع: إلغاء تصنيع القديم
        if old_factory_id and hasattr(models, "ManufacturingBatch") and (purchase.total_qty or 0) > 0:
            old_total_qty, old_total_mfg, old_avg_mfg = _calc_mfg_totals(purchase)
            if old_total_qty > 0:
                mb_rev = models.ManufacturingBatch(
                    date=purchase.date or date,
                    factory_id=old_factory_id,
                    product_id=None,
                    qty=-(int(old_total_qty)),
                    unit_cost=float(old_avg_mfg),
                    note=_po_total_note(code, purchase.id, "edit_rev"),
                )
                db.add(mb_rev)

        # ✅ 2) رجوع تأثير البنود القديمة على المخزون
        for item in purchase.items:
            product = db.query(models.Product).get(item.product_id)
            if not product:
                continue

            dec_qty = int(item.qty or 0)
            if dec_qty <= 0:
                continue

            old_stock_now = int(product.stock or 0)
            product.stock = max(0, old_stock_now - dec_qty)

            try:
                if (
                    int(product.stock + dec_qty) == int(getattr(item, "applied_stock", 0) or 0)
                    and float(product.cost_price or 0) == float(getattr(item, "applied_cost_price", 0) or 0)
                    and float(product.price or 0) == float(getattr(item, "applied_price", 0) or 0)
                ):
                    product.cost_price = float(getattr(item, "old_cost_price", 0) or 0)
                    product.price = float(getattr(item, "old_price", 0) or 0)
            except Exception:
                pass

            if hasattr(models, "ProductMovement"):
                mv = models.ProductMovement(
                    product_id=product.id,
                    movement_type="purchase_edit_revert",
                    qty=-(dec_qty),
                    balance_after=int(product.stock or 0),
                    ref_type="PurchaseInvoice",
                    ref_id=int(purchase.id),
                    note=_pm_note(code, "revert", -dec_qty),
                )
                db.add(mv)

        # امسح البنود القديمة
        for it in list(purchase.items):
            db.delete(it)

        # ✅ 3) تطبيق الشيت الجديد
        total_qty = 0
        total_cost = 0.0
        total_mfg_cost = 0.0
        total_fabric_cost = 0.0
        total_overhead_cost = 0.0

        for barcode, name, color, size, qty, unit_fabric_cost, unit_mfg_cost, unit_over_cost, unit_total_cost, unit_price in _iter_rows_any_format(ws):
            if not barcode or qty <= 0:
                continue

            unit_fabric_cost = float(unit_fabric_cost or 0)
            unit_mfg_cost = float(unit_mfg_cost or 0)
            unit_over_cost = float(unit_over_cost or 0)
            unit_total_cost = float(unit_total_cost or 0)
            unit_price = float(unit_price or 0)

            if unit_total_cost <= 0 and (unit_fabric_cost + unit_mfg_cost + unit_over_cost) > 0:
                unit_total_cost = float(unit_fabric_cost + unit_mfg_cost + unit_over_cost)

            if unit_total_cost + 1e-9 < (unit_fabric_cost + unit_mfg_cost + unit_over_cost):
                db.rollback()
                return RedirectResponse(
                    url=_back_url("⚠️ خطأ: total_cost لازم يكون >= (fabric_cost + mfg_cost + overhead_cost)"),
                    status_code=303
                )

            product = db.query(models.Product).filter(models.Product.barcode == barcode).first()
            if not product:
                product = models.Product(
                    name=name or barcode,
                    barcode=barcode,
                    color=color or None,
                    size=size or None,
                    cost_price=0,
                    price=0,
                    stock=0,
                )
                db.add(product)
                db.flush()

            prev_cost_avg = float(product.cost_price or 0)
            prev_price = float(product.price or 0)
            prev_stock = int(product.stock or 0)

            new_stock = prev_stock + int(qty)
            new_cost_avg = _weighted_avg(prev_stock, prev_cost_avg, int(qty), unit_total_cost)

            new_price = prev_price
            if unit_price > 0:
                new_price = unit_price

            product.stock = new_stock
            product.cost_price = new_cost_avg
            product.price = new_price

            item = models.PurchaseItem(
                purchase_id=purchase.id,
                product_id=product.id,
                barcode=barcode,
                name=name or product.name,
                color=color or product.color,
                size=size or product.size,
                qty=int(qty),

                unit_fabric_cost=unit_fabric_cost,
                unit_mfg_cost=unit_mfg_cost,
                unit_overhead_cost=unit_over_cost,
                unit_total_cost=unit_total_cost,

                unit_cost=unit_total_cost,
                unit_price=unit_price,

                old_cost_price=prev_cost_avg,
                old_price=prev_price,
                old_stock=prev_stock,

                applied_cost_price=new_cost_avg,
                applied_price=new_price,
                applied_stock=new_stock,
            )
            db.add(item)

            if hasattr(models, "ProductMovement"):
                mv = models.ProductMovement(
                    product_id=product.id,
                    movement_type="purchase_edit_apply",
                    qty=int(qty),
                    balance_after=int(product.stock or 0),
                    ref_type="PurchaseInvoice",
                    ref_id=int(purchase.id),
                    note=_pm_note(code, "apply", int(qty), unit_total_cost, unit_mfg_cost, unit_fabric_cost, unit_over_cost, unit_price),
                )
                db.add(mv)

            total_qty += int(qty)
            total_cost += int(qty) * unit_total_cost
            total_mfg_cost += int(qty) * unit_mfg_cost
            total_fabric_cost += int(qty) * unit_fabric_cost
            total_overhead_cost += int(qty) * unit_over_cost

        old_qty_show = purchase.total_qty or 0
        old_cost_show = purchase.total_cost or 0

        purchase.total_qty = float(total_qty)
        purchase.total_cost = float(total_cost)
        purchase.date = date
        purchase.file_name = file.filename
        if note:
            purchase.note = note

        # ✅ 4) حساب المصنع: إضافة تصنيع الجديد
        if new_factory_id and hasattr(models, "ManufacturingBatch") and total_qty > 0:
            avg_mfg = _avg_cost(int(total_qty), float(total_mfg_cost))
            mb_add = models.ManufacturingBatch(
                date=date,
                factory_id=new_factory_id,
                product_id=None,
                qty=int(total_qty),
                unit_cost=float(avg_mfg),
                note=_po_total_note(code, purchase.id, "edit_add"),
            )
            db.add(mb_add)

        # ✅ 5) خصم خامات القماش الجديدة
        if total_fabric_cost > 0:
            _materials_post(
                db,
                type_="OUT",
                amount=round(float(total_fabric_cost), 2),
                date=date,
                note=f"[PO {code}] تعديل — خصم خامات (قماش فقط) للجديد",
                ref_type="PurchaseInvoice",
                ref_id=int(purchase.id),
            )

        now = datetime.now().strftime("%Y-%m-%d %H:%M")
        h = purchase.history or ""
        h += (
            f"\n[{now}] تعديل الفاتورة من ملف جديد: {file.filename} "
            f"(الكميات {old_qty_show}→{total_qty}، إجمالي تكلفة المخزون {old_cost_show:.2f}→{total_cost:.2f}) "
            f"+ (خامات القماش اتعملها رجوع/خصم تلقائي) + (حساب المصنع سطر تصنيع إجمالي)."
        )
        purchase.history = h

        db.commit()
        return RedirectResponse(url=f"/purchases/{purchase.id}?msg=✅ تم تعديل فاتورة المشتريات", status_code=303)

    except Exception:
        db.rollback()
        return RedirectResponse(url=_back_url("⚠️ حصل خطأ أثناء تعديل الفاتورة"), status_code=303)